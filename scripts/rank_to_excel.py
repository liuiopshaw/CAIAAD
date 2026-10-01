#!/usr/bin/env python3
"""
Export a ranking run to an Excel workbook (.xlsx).

Reuses rank_designer_outputs.parse_records (same parsing/sorting as the MD ranking)
plus the formula_lookup map, and writes:
  Sheet 1 "Ranking": the full ranked table (same columns as ranking_<TS>.md)
  Sheet 2 "Run Metrics": classification distributions (Drug_Type/Modality,
    Target_Category/Disease_Intervention, AD_Mechanism/Mechanism by era),
    identifier coverage (SMILES/Target_UniProt non-NA), neutral element
    composition statistics
  Sheet 3 "Type_Top5": top-5 per drug type / modality by the primary sort key

Usage: python scripts/rank_to_excel.py [timestamp|rawfile.txt] [dedup] [--rubric path]

Sorting: like rank_designer_outputs.py, when the run directory contains
subscores_<TS>.json the deterministic Score_Adj is the primary sort key and the
designer self-report column is shown as SelfReport; otherwise the self-report
ranking is used unchanged. --rubric re-scores historical subscores without
re-running agents.
"""

import io, json, re, sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from rank_designer_outputs import (parse_records, _normalize_any, _detect_era,
                              FIELDS, LEGACY_FIELDS,
                              RECORD_LENGTHS, TABLE_FIELDS, OUTPUT,
                              SELF_SCORE, NAME, MODALITY, SMILES, UNIPROT,
                              ELEMENTS, INTERVENTION, MECHANISM, NADH,
                              display_pairs,
                              load_subscores, compute_score_map, SELF_REPORT, SCORE_ADJ)
from output_utils import find_run_dir
import output_schema

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def main():
    argv = sys.argv[1:]
    rubric_path = None
    if "--rubric" in argv:
        i = argv.index("--rubric")
        rubric_path = argv[i + 1]
        del argv[i:i + 2]
    arg = argv[0] if argv else None
    file_mode = bool(arg and arg.endswith(".txt"))
    dedup = len(argv) > 1 and argv[1].lower() == "dedup"

    if file_mode:
        # Direct raw-output file (e.g. base_task100_<TS>.txt)
        src_path = Path(arg)
        if not src_path.is_absolute():
            src_path = OUTPUT / arg
        ts = src_path.stem
        src_files = [src_path.name]
        records, fused, bad = [], [], []
        is_current = False
        text = io.open(src_path, encoding="utf-8").read()
        file_era = _detect_era(text)
        for line in text.split("\n"):
            line = line.strip()
            if not line or "|" not in line:
                continue
            cells = [c.strip() for c in line.split("|")]
            if cells[0] == "Material_Name":
                continue  # field-spec header line emitted by the designer prompt
            if len(cells) in RECORD_LENGTHS:
                rec, is_cur = _normalize_any(cells, file_era)
                is_current = is_current or is_cur
                if rec is not None:
                    records.append(rec)
                else:
                    bad.append((src_path.name, line))
            else:
                bad.append((src_path.name, line))
    else:
        ts = arg
        if ts is None:
            candidates = (list(OUTPUT.glob("task100_designer_*_part1.txt"))
                          + list(OUTPUT.glob("run_*/task100_designer_*_part1.txt")))
            if not candidates:
                raise SystemExit("No task100_designer_*_part1.txt found under outputs/ — run the pipeline first")
            latest = max(candidates, key=lambda p: p.stat().st_mtime)
            ts = re.search(r"task100_designer_(\d+)_part1", latest.name).group(1)
        src_files, records, fused, bad, is_current = parse_records(ts)

    def score(rec):
        try:
            return float(rec[SELF_SCORE])
        except ValueError:
            return float("nan")

    valid = [r for r in records if score(r) == score(r)]

    # Phase 3: deterministic score from extracted subscores (run mode only).
    rubric, score_map = (None, {})
    if not file_mode:
        sub_payload = load_subscores(ts)
        if sub_payload is not None:
            rubric, score_map = compute_score_map(sub_payload, rubric_path)

    if score_map:
        valid.sort(key=lambda r: (-score_map[r[NAME]]["total_adj"], r[NAME])
                   if r[NAME] in score_map else (float("inf"), r[NAME]))
    else:
        valid.sort(key=lambda r: (-score(r), r[NAME]))

    # Optional: deduplicate by Material_Name (keep first = highest score after
    # sorting). Requested for the base-model control output, whose template
    # loop emitted many identical materials. Raw outputs are never modified.
    removed = 0
    if dedup:
        seen = set()
        uniq = []
        for r in valid:
            if r[NAME] not in seen:
                seen.add(r[NAME])
                uniq.append(r)
        removed = len(valid) - len(uniq)
        valid = uniq

    # DB-verified formulas (tool data)
    db_map, db_source = {}, None
    map_dir = src_path.parent if file_mode else find_run_dir(ts)
    map_file = map_dir / f"task100_formula_map_{ts}.json"
    if map_file.exists():
        payload = json.loads(map_file.read_text(encoding="utf-8"))
        db_map = payload.get("materials", {})
        db_source = payload.get("source")

    wb = Workbook()

    # ---- Sheet 1: Ranking ----
    ws = wb.active
    ws.title = "Ranking"
    pairs = display_pairs(is_current)
    # Deterministic-scoring mode: rename self-report column, insert Score_Adj before it.
    base_pairs = []
    for h, k in pairs:
        if score_map and k == SELF_SCORE:
            base_pairs += [(SCORE_ADJ, "__adj__"), (SELF_REPORT, SELF_SCORE)]
        else:
            base_pairs.append((h, k))
    db_pos = next(i for i, (h, k) in enumerate(base_pairs) if k == "Chemical_Formula") + 1
    final_pairs = base_pairs[:db_pos] + ([("DB_Formula", "__db__")] if db_map else []) + base_pairs[db_pos:]
    header = ["Rank"] + [h for h, k in final_pairs]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for i, r in enumerate(valid, 1):
        cells = []
        for h, k in final_pairs:
            if k == "__db__":
                cells.append((db_map.get(r[NAME]) or {}).get("combined_formula") or "—")
            elif k == "__adj__":
                res = score_map.get(r[NAME])
                cells.append(round(res["total_adj"], 3) if res else "—")
            elif score_map and k == SELF_SCORE:
                cells.append(r[SELF_SCORE])
            else:
                cells.append(r[k])
        row = [i] + cells
        self_col = next(j for j, (h, k) in enumerate(final_pairs) if h == (SELF_REPORT if score_map else SELF_SCORE)) + 1  # +1 for Rank
        try:
            row[self_col] = float(row[self_col])
        except (ValueError, TypeError):
            pass
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(valid)+1}"
    for idx, width in enumerate([6, 24, 14, 16, 24, 18, 16, 14, 14, 10, 24, 24, 12, 60], 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    # ---- Sheet 2: Run Metrics ----
    ws2 = wb.create_sheet("Run Metrics")

    mod_freq = Counter(r[MODALITY].strip() for r in valid)
    cat_freq = Counter(r[INTERVENTION].strip() for r in valid)
    mech_freq = Counter(r[MECHANISM].strip() for r in valid)
    smiles_ok = sum(1 for r in valid if r[SMILES].strip() not in ("", "NA"))
    uniprot_ok = sum(1 for r in valid if r[UNIPROT].strip() not in ("", "NA"))

    all_elem = Counter()
    for r in valid:
        for e in r[ELEMENTS].split(","):
            e = e.strip()
            if e and e not in ("O", "N", "C"):
                all_elem[e] += 1

    type_label = "Drug_Type" if is_current else "Modality"
    type_enums = output_schema.DRUG_TYPES if is_current else output_schema.MODALITIES
    rows = [
        ("Run Timestamp", ts),
        ("Source Files", ", ".join(src_files)),
        ("Scoring Mode", "Score_Adj deterministic computation (scoring_rubric)" if score_map else "designer self-reported score (no subscores, fallback)"),
        ("DB_Formula Source", db_source or "not verified"),
        ("Total Materials", f"{len(valid)}" + (f" (deduplicated by Material_Name: {len(valid)+removed} original rows, {removed} duplicates removed)" if dedup else "")),
        ("Score Range", f"{score(valid[-1])} - {score(valid[0])}" if valid else ""),
        ("", ""),
        (f"{type_label} Distribution", ", ".join(f"{m}×{mod_freq.get(m, 0)}" for m in type_enums)),
        ("Target_Category Distribution" if is_current else "Disease_Intervention Distribution", ", ".join(f"{k}×{v}" for k, v in cat_freq.most_common())),
        ("AD_Mechanism Distribution" if is_current else "Mechanism Distribution", ", ".join(f"{k}×{v}" for k, v in mech_freq.most_common())),
        ("Element frequency Top8, full ranking (excl. O/N/C)", ", ".join(f"{e}×{n}" for e, n in all_elem.most_common(8))),
        ("", ""),
        ("SMILES non-NA", smiles_ok),
        ("Target_UniProt non-NA", uniprot_ok),
    ]
    for k, v in rows:
        ws2.append([k, v])
    ws2.column_dimensions["A"].width = 36
    ws2.column_dimensions["B"].width = 70
    for row in ws2.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.font = Font(bold=True)

    # ---- Sheet 3: Type_Top5 (primary sort key — `valid` is already sorted) ----
    ws3 = wb.create_sheet("Type_Top5")
    ws3.append([type_label, "#", "Material_Name", SCORE_ADJ if score_map else SELF_SCORE])
    for cell in ws3[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    observed = list(type_enums) + sorted(
        {r[MODALITY].strip() for r in valid} - set(type_enums))
    for m in observed:
        group = [r for r in valid if r[MODALITY].strip() == m][:5]
        for j, r in enumerate(group, 1):
            if score_map:
                res = score_map.get(r[NAME])
                score = round(res["total_adj"], 3) if res else "—"
            else:
                score = r[SELF_SCORE]
                try:
                    score = float(score)
                except (ValueError, TypeError):
                    pass
            ws3.append([m, j, r[NAME], score])
    for col, width in zip("ABCD", (16, 4, 30, 14)):
        ws3.column_dimensions[col].width = width

    xlsx_path = map_dir / ((f"{ts}_dedup.xlsx" if dedup else f"{ts}.xlsx") if file_mode else f"ranking_{ts}.xlsx")
    wb.save(xlsx_path)
    print(f"Wrote {xlsx_path} ({len(valid)} rows + Run Metrics/Type_Top5 sheets, {len(bad)} unparsed)")


if __name__ == "__main__":
    main()
